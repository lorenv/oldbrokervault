#!/usr/bin/env python3
"""
SDE Analyzer - Standalone Module
Analyzes P&L statements to calculate Seller's Discretionary Earnings (SDE)

Features:
- Detects and removes empty columns
- Aggregates monthly data to annual
- Identifies revenue, NOI, and add-backs
- Creates professionally formatted Excel output with formulas
- Preserves all original source sheets

Author: Generated for standalone use
License: MIT
"""

import pandas as pd
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import re
import os
import json
from typing import Dict, List, Tuple, Optional


class SDEAnalyzer:
    """Main class for analyzing P&L statements and generating SDE reports"""
    
    def __init__(self, input_file: str, output_file: str, company_name: Optional[str] = None):
        """
        Initialize the analyzer
        
        Args:
            input_file: Path to input Excel file
            output_file: Path for output Excel file
            company_name: Optional company name (extracted from filename if not provided)
        """
        self.input_file = input_file
        self.output_file = output_file
        self.company_name = company_name or self._extract_company_name()
        
        self.df_cleaned = None
        self.source_sheet_name = None
        self.col_mapping = {}
        self.structure = {}
        self.addbacks = []
    
    def _extract_company_name(self) -> str:
        """Extract company name from filename"""
        filename = os.path.basename(self.input_file)
        company_name = filename.replace('.xlsx', '').replace('_', ' ').replace('-', ' ')
        
        # Remove common P&L terms
        exclude_terms = ['profit', 'loss', 'and', 'p&l', 'pl', 'monthly', 'annual']
        parts = company_name.split()
        company_name = ' '.join([p for p in parts if p.lower() not in exclude_terms])
        
        return company_name.strip() or "Company"
    
    def detect_empty_columns(self) -> Tuple[pd.DataFrame, str, Dict]:
        """
        Detect and track empty columns
        
        Returns:
            Tuple of (cleaned_dataframe, sheet_name, column_mapping)
        """
        wb = load_workbook(self.input_file, data_only=True)
        sheet_name = wb.sheetnames[0]

        df = pd.read_excel(self.input_file, sheet_name=sheet_name, header=None, engine='openpyxl')
        
        # Identify empty columns
        empty_cols = []
        for col_idx in range(df.shape[1]):
            col_data = df.iloc[:, col_idx]
            is_empty = (
                col_data.isna().all() or 
                (col_data.fillna('').astype(str).str.strip() == '').all()
            )
            if is_empty:
                empty_cols.append(col_idx)
        
        # Create mapping of cleaned index to original index
        col_mapping = {}
        cleaned_idx = 0
        for orig_idx in range(df.shape[1]):
            if orig_idx not in empty_cols:
                col_mapping[cleaned_idx] = orig_idx
                cleaned_idx += 1
        
        # Remove empty columns
        df_cleaned = df.drop(columns=empty_cols) if empty_cols else df
        
        return df_cleaned, sheet_name, col_mapping
    
    def identify_structure(self, df: pd.DataFrame, col_mapping: Dict) -> Dict:
        """
        Identify financial structure (years, revenue row, NOI row)
        
        Args:
            df: Cleaned dataframe
            col_mapping: Mapping of cleaned to original column indices
            
        Returns:
            Dictionary with structure information
        """
        month_cols_by_year = {}
        month_year_pattern = re.compile(
            r'(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\s*[\'"]?(\d{2,4})',
            re.IGNORECASE
        )
        
        # Find month/year columns
        for row_idx in range(min(10, len(df))):
            for col_idx in range(df.shape[1]):
                cell_value = str(df.iloc[row_idx, col_idx]).strip()
                
                if cell_value == 'nan' or not cell_value:
                    continue
                
                match = month_year_pattern.search(cell_value)
                if match:
                    year_str = match.group(2)
                    year = '20' + year_str if len(year_str) == 2 else year_str
                    
                    orig_col = col_mapping[col_idx]
                    
                    if year not in month_cols_by_year:
                        month_cols_by_year[year] = []
                    month_cols_by_year[year].append(orig_col)
        
        # Find Revenue row
        revenue_row = None
        for row_idx in range(len(df)):
            for check_col in range(min(5, df.shape[1])):
                check_label = str(df.iloc[row_idx, check_col]).lower() if pd.notna(df.iloc[row_idx, check_col]) else ""
                if any(term in check_label for term in ['total income', 'total revenue']):
                    revenue_row = row_idx
                    break
            if revenue_row:
                break
        
        # Find NOI row
        noi_row = None
        for row_idx in range(len(df)):
            for check_col in range(min(5, df.shape[1])):
                check_label = str(df.iloc[row_idx, check_col]).lower() if pd.notna(df.iloc[row_idx, check_col]) else ""
                if any(term in check_label for term in ['net income', 'net operating income']):
                    noi_row = row_idx
                    break
            if noi_row:
                break
        
        return {
            'revenue_row': revenue_row,
            'noi_row': noi_row,
            'month_cols_by_year': month_cols_by_year
        }
    
    def detect_addbacks(self, df: pd.DataFrame, structure: Dict) -> List[Dict]:
        """
        Detect add-back line items between revenue and NOI

        Args:
            df: Cleaned dataframe
            structure: Structure dictionary

        Returns:
            List of add-back dictionaries
        """
        # Check for AI-generated analysis JSON file
        json_path = self.input_file.replace('.xlsx', '_ai_analysis.json').replace('.xls', '_ai_analysis.json')
        if os.path.exists(json_path):
            try:
                with open(json_path, 'r') as f:
                    ai_result = json.load(f)

                print(f"\n🤖 AI Analysis Results:")

                # Update revenue row if AI found it
                if ai_result.get('revenue_row'):
                    old_revenue = structure.get('revenue_row')
                    structure['revenue_row'] = ai_result['revenue_row'] - 1  # Convert to 0-based
                    print(f"   📈 Revenue row: {ai_result['revenue_row']} (AI detected)")
                    if old_revenue != structure['revenue_row']:
                        print(f"      (Overriding pattern match: {old_revenue + 1 if old_revenue else 'none'})")

                # Update NOI row if AI found it
                if ai_result.get('noi_row'):
                    old_noi = structure.get('noi_row')
                    structure['noi_row'] = ai_result['noi_row'] - 1  # Convert to 0-based
                    print(f"   📊 NOI row: {ai_result['noi_row']} (AI detected)")
                    if old_noi != structure['noi_row']:
                        print(f"      (Overriding pattern match: {old_noi + 1 if old_noi else 'none'})")

                # Use AI-detected add-backs
                if ai_result.get('addbacks'):
                    print(f"   ✓ Add-backs: {len(ai_result['addbacks'])} identified by AI")
                    addbacks = []
                    for ab in ai_result['addbacks']:
                        # Convert AI row (1-based) to 0-based index
                        row_idx = ab['row'] - 1
                        if 0 <= row_idx < df.shape[0]:
                            addbacks.append({
                                'row': row_idx,
                                'label': ab['label'],
                                'category': ab.get('category', 'ai_detected')
                            })
                            print(f"      - {ab['label']} (row {ab['row']})")
                    if addbacks:
                        return addbacks
                    else:
                        print(f"   ⚠ No valid add-backs after filtering")
            except Exception as e:
                print(f"⚠ Failed to load AI analysis: {e}, falling back to pattern matching")

        # Fall back to original pattern-based detection
        addback_patterns = {
            'depreciation': ['depreciation', 'deprec', 'amortization'],
            'owner_comp': ['owner compensation', 'owner salary', 'owner wages'],
            'officer_comp': ['officer compensation', 'officer salary'],
            'interest': ['interest expense', 'interest paid'],
            'auto': ['auto expense', 'vehicle expense', 'car expense'],
            'travel': ['travel expense', 'meals', 'entertainment'],
            'professional': ['legal', 'accounting', 'professional fees'],
            'insurance': ['health insurance', 'life insurance'],
            'payroll_tax': ['payroll tax', '941', 'futa', 'suta'],
            'bonus': ['bonus']
        }
        
        addbacks = []
        
        if structure['revenue_row'] is None or structure['noi_row'] is None:
            return addbacks
        
        # Scan between revenue and NOI
        for row_idx in range(structure['revenue_row'] + 1, structure['noi_row']):
            label = ""
            label_text = ""
            
            for col_idx in range(min(5, df.shape[1])):
                cell_val = str(df.iloc[row_idx, col_idx]) if pd.notna(df.iloc[row_idx, col_idx]) else ""
                if cell_val and cell_val != 'nan':
                    label = cell_val.lower().strip()
                    label_text = df.iloc[row_idx, col_idx]
                    break
            
            if not label:
                continue
            
            for category, patterns in addback_patterns.items():
                if any(pattern in label for pattern in patterns):
                    addbacks.append({
                        'row': row_idx,
                        'label': label_text,
                        'category': category
                    })
                    break
        
        return addbacks
    
    def create_output(self):
        """Create the professionally formatted SDE output file"""
        # Load original file
        wb_source = load_workbook(self.input_file, data_only=False)
        wb_out = Workbook()
        wb_out.remove(wb_out.active)
        
        # Copy ALL original sheets
        for sheet_name in wb_source.sheetnames:
            source_sheet = wb_source[sheet_name]
            dest_sheet = wb_out.create_sheet(sheet_name)
            
            for row in source_sheet.iter_rows():
                for cell in row:
                    dest_cell = dest_sheet[cell.coordinate]
                    if cell.value:
                        dest_cell.value = cell.value
                    if cell.has_style:
                        if cell.font:
                            dest_cell.font = cell.font.copy()
                        if cell.number_format:
                            dest_cell.number_format = cell.number_format
                        if cell.fill:
                            dest_cell.fill = cell.fill.copy()
        
        # Create SDE Analysis sheet as FIRST sheet
        ws = wb_out.create_sheet("SDE Analysis", 0)  # Insert at index 0 (first position)
        
        # Apply styling
        self._apply_styling(ws)
        
        wb_out.save(self.output_file)
        wb_out.close()
    
    def _apply_styling(self, ws):
        """Apply professional styling to the SDE Analysis worksheet"""
        # Style definitions
        title_font = Font(bold=True, size=14)
        subtitle_font = Font(bold=True, size=12)
        date_font = Font(size=11)
        header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
        header_font = Font(bold=True, size=12, color='FFFFFF')
        section_fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
        section_font = Font(bold=True, size=11)
        subtotal_fill = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')
        bold_font = Font(bold=True, size=11)
        normal_font = Font(size=11)
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        center_aligned = Alignment(horizontal='center', vertical='center')
        
        # Get years
        years = sorted(self.structure['month_cols_by_year'].keys())
        if years:
            year_range = f"January 1, {years[0]} - December 31, {years[-1]}" if len(years) > 1 else f"Year {years[0]}"
        else:
            year_range = "Financial Analysis"
        
        # Title section (rows 1-4)
        ws.merge_cells('A1:E1')
        ws['A1'] = "SELLER'S DISCRETIONARY EARNINGS (SDE) ANALYSIS"
        ws['A1'].font = title_font
        ws['A1'].alignment = center_aligned
        
        ws.merge_cells('A2:E2')
        ws['A2'] = self.company_name
        ws['A2'].font = subtitle_font
        ws['A2'].alignment = center_aligned
        
        ws.merge_cells('A3:E3')
        ws['A3'] = year_range
        ws['A3'].font = date_font
        ws['A3'].alignment = center_aligned
        
        # Column widths
        ws.column_dimensions['A'].width = 40
        for idx in range(2, 2 + len(years)):
            ws.column_dimensions[get_column_letter(idx)].width = 15
        
        # Header row (row 5)
        row = 5
        ws['A5'] = "Line Item"
        ws['A5'].fill = header_fill
        ws['A5'].font = header_font
        ws['A5'].alignment = center_aligned
        ws['A5'].border = thin_border
        
        for idx, year in enumerate(years, start=2):
            col_letter = get_column_letter(idx)
            ws[f'{col_letter}{row}'] = year
            ws[f'{col_letter}{row}'].fill = header_fill
            ws[f'{col_letter}{row}'].font = header_font
            ws[f'{col_letter}{row}'].alignment = center_aligned
            ws[f'{col_letter}{row}'].border = thin_border
        
        # NET INCOME CALCULATION section
        row = 6
        ws[f'A{row}'] = "NET INCOME CALCULATION"
        ws[f'A{row}'].fill = section_fill
        ws[f'A{row}'].font = section_font
        
        # Revenue row
        row = 7
        revenue_row_num = row
        ws[f'A{row}'] = "Total Revenue"
        ws[f'A{row}'].font = normal_font
        if self.structure['revenue_row'] is not None:
            for idx, year in enumerate(years, start=2):
                orig_cols = self.structure['month_cols_by_year'][year]
                col_refs = [f"'{self.source_sheet_name}'!{get_column_letter(c+1)}{self.structure['revenue_row']+1}" for c in orig_cols]
                formula = f"=SUM({','.join(col_refs)})"
                ws.cell(row, idx).value = formula
                ws.cell(row, idx).number_format = '$#,##0'
                ws.cell(row, idx).font = normal_font
        
        # NOI row
        row = 8
        noi_row_num = row
        ws[f'A{row}'] = "Net Operating Income"
        ws[f'A{row}'].font = bold_font
        ws[f'A{row}'].fill = subtotal_fill
        if self.structure['noi_row'] is not None:
            for idx, year in enumerate(years, start=2):
                orig_cols = self.structure['month_cols_by_year'][year]
                col_refs = [f"'{self.source_sheet_name}'!{get_column_letter(c+1)}{self.structure['noi_row']+1}" for c in orig_cols]
                formula = f"=SUM({','.join(col_refs)})"
                ws.cell(row, idx).value = formula
                ws.cell(row, idx).number_format = '$#,##0'
                ws.cell(row, idx).font = bold_font
                ws.cell(row, idx).fill = subtotal_fill
        
        # ADD-BACKS section
        row = 10
        ws[f'A{row}'] = "ADD-BACKS"
        ws[f'A{row}'].fill = section_fill
        ws[f'A{row}'].font = section_font
        row += 1
        
        addback_start_row = row
        for addback in self.addbacks:
            ws.cell(row, 1).value = addback['label']
            ws.cell(row, 1).font = normal_font
            for idx, year in enumerate(years, start=2):
                orig_cols = self.structure['month_cols_by_year'][year]
                col_refs = [f"'{self.source_sheet_name}'!{get_column_letter(c+1)}{addback['row']+1}" for c in orig_cols]
                formula = f"=SUM({','.join(col_refs)})"
                ws.cell(row, idx).value = formula
                ws.cell(row, idx).number_format = '$#,##0'
                ws.cell(row, idx).font = normal_font
            row += 1
        
        addback_end_row = row - 1
        
        # Total add-backs
        row += 1
        total_addbacks_row = row
        ws[f'A{row}'] = "Total Add-Backs"
        ws[f'A{row}'].font = bold_font
        ws[f'A{row}'].fill = subtotal_fill
        for idx, year in enumerate(years, start=2):
            col_letter = get_column_letter(idx)
            if self.addbacks:
                formula = f"=SUM({col_letter}{addback_start_row}:{col_letter}{addback_end_row})"
            else:
                formula = "=0"
            ws[f'{col_letter}{row}'] = formula
            ws[f'{col_letter}{row}'].number_format = '$#,##0'
            ws[f'{col_letter}{row}'].font = bold_font
            ws[f'{col_letter}{row}'].fill = subtotal_fill
        
        # TOTAL SDE
        row += 2
        ws[f'A{row}'] = "TOTAL SDE (NOI + Add-Backs)"
        ws[f'A{row}'].font = Font(bold=True, size=12)
        ws[f'A{row}'].fill = subtotal_fill
        for idx, year in enumerate(years, start=2):
            col_letter = get_column_letter(idx)
            formula = f"={col_letter}{noi_row_num}+{col_letter}{total_addbacks_row}"
            ws[f'{col_letter}{row}'] = formula
            ws[f'{col_letter}{row}'].number_format = '$#,##0'
            ws[f'{col_letter}{row}'].font = Font(bold=True, size=12)
            ws[f'{col_letter}{row}'].fill = subtotal_fill
    
    def analyze(self, verbose: bool = True) -> Dict:
        """
        Run the complete analysis
        
        Args:
            verbose: Whether to print progress messages
            
        Returns:
            Dictionary with analysis results
        """
        if verbose:
            print(f"\n{'='*60}")
            print("SDE ANALYSIS")
            print(f"{'='*60}")
            print(f"Input: {self.input_file}")
            print(f"Output: {self.output_file}")
            print(f"Company: {self.company_name}")
        
        # Step 1: Detect empty columns
        if verbose:
            print(f"\n{'='*60}")
            print("STEP 1: Detecting empty columns...")
        self.df_cleaned, self.source_sheet_name, self.col_mapping = self.detect_empty_columns()
        
        # Step 2: Identify structure
        if verbose:
            print("STEP 2: Identifying financial structure...")
        self.structure = self.identify_structure(self.df_cleaned, self.col_mapping)
        
        # Step 3: Detect add-backs
        if verbose:
            print("STEP 3: Detecting add-backs...")
        self.addbacks = self.detect_addbacks(self.df_cleaned, self.structure)
        if verbose and self.addbacks:
            for ab in self.addbacks:
                print(f"  - {ab['label']}")
        
        # Step 4: Create output
        if verbose:
            print("STEP 4: Creating output file...")
        self.create_output()
        
        if verbose:
            print(f"\n{'='*60}")
            print("✓ Analysis complete!")
            print(f"{'='*60}")
        
        return {
            'years': sorted(self.structure['month_cols_by_year'].keys()),
            'addbacks_count': len(self.addbacks),
            'output_file': self.output_file
        }


def analyze_sde(input_file: str, output_file: str, company_name: Optional[str] = None, verbose: bool = True) -> Dict:
    """
    Convenience function to analyze a P&L and generate SDE report
    
    Args:
        input_file: Path to input Excel file
        output_file: Path for output Excel file
        company_name: Optional company name
        verbose: Whether to print progress
        
    Returns:
        Dictionary with analysis results
    """
    analyzer = SDEAnalyzer(input_file, output_file, company_name)
    return analyzer.analyze(verbose=verbose)


if __name__ == "__main__":
    import sys
    
    if len(sys.argv) < 3:
        print("Usage: python sde_analyzer.py <input.xlsx> <output.xlsx> [company_name]")
        sys.exit(1)
    
    input_file = sys.argv[1]
    output_file = sys.argv[2]
    company_name = sys.argv[3] if len(sys.argv) > 3 else None
    
    try:
        result = analyze_sde(input_file, output_file, company_name)
        print(f"\nSuccess! Generated SDE analysis for {len(result['years'])} years")
        print(f"Output: {result['output_file']}")
    except Exception as e:
        print(f"\nError: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)
